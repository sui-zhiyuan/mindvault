#!/usr/bin/env python3
"""pptx-read / dump_pptx.py

Structured, read-only dump of a .pptx through python-pptx.

This complements the PowerPoint-COM path: COM gives a pixel-perfect picture of
each slide, this gives exact structure -- shape names and types, geometry in
points, tables cell by cell, and speaker notes -- which is what you want when
the text has to be reasoned over rather than looked at.

Run it through uv so nothing is installed into the system interpreter:

    UV_CACHE_DIR=<writable-dir> uv run --no-project --with python-pptx -- \
        python dump_pptx.py <deck.pptx> [--format md|json] [--out FILE]
        [--notes/--no-notes] [--geometry/--no-geometry]
"""

from __future__ import annotations

import argparse
import json
import sys

from pptx import Presentation
from pptx.util import Emu

EMU_PER_PT = 12700


def pt(v) -> float | None:
    """EMU -> points, rounded; None when the shape has no such value."""
    if v is None:
        return None
    return round(Emu(v).pt, 1)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text_frame.text.strip()


def table_rows(shape) -> list[list[str]]:
    if not getattr(shape, "has_table", False):
        return []
    rows = []
    for r in shape.table.rows:
        rows.append([c.text.strip() for c in r.cells])
    return rows


def collect(deck_path: str, notes: bool, geometry: bool) -> dict:
    prs = Presentation(deck_path)
    slides = []

    for idx, slide in enumerate(prs.slides, start=1):
        entry = {
            "index": idx,
            "layout": slide.slide_layout.name,
            "shapes": [],
        }
        if notes and slide.has_notes_slide:
            entry["notes"] = slide.notes_slide.notes_text_frame.text.strip()

        for shape in slide.shapes:
            item = {
                "name": shape.name,
                "type": str(shape.shape_type).split(" ")[0] if shape.shape_type else "",
            }
            text = shape_text(shape)
            if text:
                item["text"] = text
            rows = table_rows(shape)
            if rows:
                item["table"] = rows
            if geometry:
                item["box_pt"] = {
                    "left": pt(shape.left),
                    "top": pt(shape.top),
                    "width": pt(shape.width),
                    "height": pt(shape.height),
                }
            entry["shapes"].append(item)

        slides.append(entry)

    return {
        "file": deck_path,
        "slides": len(slides),
        "size_pt": {
            "width": round(Emu(prs.slide_width).pt, 1),
            "height": round(Emu(prs.slide_height).pt, 1),
        },
        "aspect": (
            round(prs.slide_width / prs.slide_height, 4) if prs.slide_height else None
        ),
        "deck": slides,
    }


def to_markdown(d: dict) -> str:
    out = [
        f"# {d['file'].split('/')[-1]}",
        "",
        f"- slides: {d['slides']}",
        f"- page size (pt): {d['size_pt']['width']} x {d['size_pt']['height']}",
        f"- aspect: {d['aspect']}",
        "",
    ]
    for s in d["deck"]:
        out.append(f"## Slide {s['index']}  (layout: {s['layout']})")
        out.append("")
        for sh in s["shapes"]:
            head = sh["name"]
            if sh.get("type"):
                head += f"  [{sh['type']}]"
            out.append(f"### {head}")
            if sh.get("box_pt"):
                b = sh["box_pt"]
                out.append(
                    f"`{b['left']},{b['top']}  {b['width']}x{b['height']} pt`"
                )
            if sh.get("text"):
                out.append(sh["text"])
            for row in sh.get("table") or []:
                out.append("| " + " | ".join(row) + " |")
            out.append("")
        if s.get("notes"):
            out.append(f"> notes: {s['notes']}")
            out.append("")
    return "\n".join(out)


def main() -> int:
    ap = argparse.ArgumentParser(description="Read-only structured dump of a .pptx")
    ap.add_argument("deck")
    ap.add_argument("--format", choices=("md", "json"), default="md")
    ap.add_argument("--out", default=None)
    ap.add_argument("--notes", dest="notes", action="store_true", default=True)
    ap.add_argument("--no-notes", dest="notes", action="store_false")
    ap.add_argument("--geometry", dest="geometry", action="store_true", default=False)
    ap.add_argument("--no-geometry", dest="geometry", action="store_false")
    args = ap.parse_args()

    data = collect(args.deck, args.notes, args.geometry)
    text = (
        json.dumps(data, ensure_ascii=False, indent=2)
        if args.format == "json"
        else to_markdown(data)
    )

    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write(text)
        print(f"wrote {args.out} ({len(text)} chars)")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
